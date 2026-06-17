#!/usr/bin/env python3
"""Validate DOCX/PDF manual deliverables produced from video-understanding work."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any


def resolve_path(value: Path) -> Path:
    if value.exists():
        return value.resolve()
    text = str(value)
    if any(ch in text for ch in "*?[]"):
        parent = value.parent if str(value.parent) else Path(".")
        matches = sorted(parent.glob(value.name))
        if matches:
            return matches[0].resolve()
    raise FileNotFoundError(value)


def read_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            table_text.extend(cell.text for cell in row.cells)
    text = "\n".join(paragraphs + table_text)
    return {
        "exists": path.exists(),
        "size_bytes": path.stat().st_size,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "text": text,
    }


def read_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return {
        "exists": path.exists(),
        "size_bytes": path.stat().st_size,
        "pages": len(reader.pages),
        "encrypted": reader.is_encrypted,
        "text": text,
    }


def keyword_hits(text: str, keywords: list[str]) -> dict[str, bool]:
    lower = text.lower()
    return {keyword: keyword.lower() in lower for keyword in keywords}


def load_keywords(inline: str, keyword_file: Path | None) -> list[str]:
    keywords = [item.strip() for item in inline.split(",") if item.strip()]
    if keyword_file:
        content = keyword_file.read_text(encoding="utf-8")
        for line in content.splitlines():
            line = line.strip()
            if line:
                keywords.append(line)
    return keywords


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="Check DOCX/PDF manual structure, sizes, and keywords.")
    parser.add_argument("--docx", type=Path, default=None)
    parser.add_argument("--pdf", type=Path, default=None)
    parser.add_argument("--keywords", default="", help="Comma-separated keywords expected in the deliverables.")
    parser.add_argument("--keywords-file", type=Path, default=None, help="UTF-8 text file with one expected keyword per line.")
    parser.add_argument("--min-size", type=int, default=50_000)
    parser.add_argument("--min-docx-images", type=int, default=1)
    parser.add_argument("--min-pdf-pages", type=int, default=1)
    args = parser.parse_args()

    keywords = load_keywords(args.keywords, args.keywords_file)
    result: dict[str, Any] = {"ok": True, "checks": {}, "errors": []}

    if args.docx:
        docx = resolve_path(args.docx)
        info = read_docx(docx)
        hits = keyword_hits(info.pop("text"), keywords)
        info["keyword_hits"] = hits
        result["checks"]["docx"] = info
        if info["size_bytes"] < args.min_size:
            result["ok"] = False
            result["errors"].append(f"DOCX is smaller than {args.min_size} bytes")
        if info["images"] < args.min_docx_images:
            result["ok"] = False
            result["errors"].append(f"DOCX has fewer than {args.min_docx_images} images")
        missing = [kw for kw, present in hits.items() if not present]
        if missing:
            result["ok"] = False
            result["errors"].append(f"DOCX missing keywords: {', '.join(missing)}")

    if args.pdf:
        pdf = resolve_path(args.pdf)
        info = read_pdf(pdf)
        hits = keyword_hits(info.pop("text"), keywords)
        info["keyword_hits"] = hits
        result["checks"]["pdf"] = info
        if info["size_bytes"] < args.min_size:
            result["ok"] = False
            result["errors"].append(f"PDF is smaller than {args.min_size} bytes")
        if info["pages"] < args.min_pdf_pages:
            result["ok"] = False
            result["errors"].append(f"PDF has fewer than {args.min_pdf_pages} pages")
        if info["encrypted"]:
            result["ok"] = False
            result["errors"].append("PDF is encrypted")
        missing = [kw for kw, present in hits.items() if not present]
        if missing:
            result["ok"] = False
            result["errors"].append(f"PDF missing keywords: {', '.join(missing)}")

    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["ok"] else 1)


if __name__ == "__main__":
    main()
